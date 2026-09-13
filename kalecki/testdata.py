"""Synthetic workbook in the exact shape of the real ledger. Every value is random; nothing
here derives from real data. `python -m kalecki.testdata out.xlsx --units 12 --months 18`."""
from __future__ import annotations

import argparse
import datetime as dt
import random

import openpyxl

LEDGER_HEADER = [None, "Lokal", "Rok", "Zarzad", "Miasto", "Adres", "nr", "wl", "mc",
                 "Czynsz zgodnie z umową", "Zaliczka na media zgodnie z umową", "Wpłata Najemcy",
                 "Faktura za zarządzanie", "Opłata do Wspólnoty", "Faktura PGE", "Naprawy z przychodu",
                 "Przelew", "Przelew spodziewany", "nadplata(+), brakuje(-)",
                 "Koszty poza podatkiem (nieoplacone przez najemce, rozliczenia mediow,etc)",
                 "Zysk-do-podatku", "data kursu", "kurs NBP"]
REFERENCE_HEADER = [None, "mieszkanie", "typ", "wartosc T0-najem", "powierzchnia", "pietro", "KW", "T0"]
TRADES_HEADER = ["Date", "Ticker", "Side", "Qty", "Price", "Fee", "Currency", "Account"]

CITIES = ["Krakow", "Warszawa", "Gdansk", "London"]
STREETS = ["Lipowa", "Dluga", "Krotka", "Polna", "Zielona", "Slowackiego", "Baker Street", "Mila"]
MANAGERS = ["-", "ZarzadCo", "MieszkaniaPlus", "HomeAdmin"]


def unit_key(city, address, nr):
    return f"{city}, {address[:15]}/ {nr}"


def make_units(rng, n):
    units = []
    for i in range(n):
        foreign = i == n - 1 and n > 2
        garage = i % 4 == 3 and not foreign
        city = "London" if foreign else rng.choice(CITIES[:3])
        address = f"{rng.choice(STREETS)} {rng.randint(1, 60)}"
        nr = f"g{rng.randint(1, 9)}" if garage else (f"{rng.randint(1, 50)}MP" if i % 7 == 6 else str(rng.randint(1, 80)))
        units.append({
            "key": unit_key(city, address, nr), "city": city, "address": address, "nr": nr,
            "kind": "garaz" if garage else "mieszkanie",
            "area": None if garage else round(rng.uniform(25, 75), 1),
            "value_t0": rng.randrange(60_000, 120_000, 5000) if garage else rng.randrange(300_000, 900_000, 10_000),
            "floor": None if garage else rng.randint(0, 6),
            "t0": dt.date(rng.randint(2010, 2022), rng.randint(1, 12), rng.randint(1, 28)),
            "manager": rng.choice(MANAGERS), "foreign": foreign,
            "rent": rng.randrange(200, 500, 50) if garage else (round(rng.uniform(900, 1500), 2) if foreign else rng.randrange(1800, 4200, 100)),
            "media": None if garage else rng.randrange(300, 900, 50),
            "mgmt": 0 if rng.random() < 0.3 else rng.choice([110, 120, 150, 180]),
            "hoa": None if garage else round(rng.uniform(300, 800), 2),
            "label": f"{city[:3]}, {address} {'parking' if garage else 'm.' + nr}",
        })
    return units


def make_workbook(n_units=12, n_months=18, seed=7, start=(2024, 1)):
    rng = random.Random(seed)
    units = make_units(rng, n_units)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "All-Data"
    ws.append([None] * len(LEDGER_HEADER))  # the real sheet has a blank first row
    ws.append(LEDGER_HEADER)
    y, m = start
    months = []
    for _ in range(n_months):
        months.append((y, m))
        m += 1
        if m > 12:
            y, m = y + 1, 1
    arrears_unit = units[1]["key"]
    vacant_unit = units[2]["key"]
    gap_unit = units[4]["key"] if n_units > 4 else None
    rows = 0
    for u in units:
        for k, (yy, mm) in enumerate(months):
            if u["key"] == gap_unit and k == 5:
                continue  # a missing month
            vacant = u["key"] == vacant_unit and k >= n_months - 3
            rent = 0 if vacant else u["rent"]
            if k >= 12 and not vacant and u["kind"] == "mieszkanie" and u["key"] != units[0]["key"]:
                rent = round(u["rent"] * 1.05, 2)  # indexed after a year; units[0] never indexed
            repairs = round(rng.uniform(50, 400), 2) if rng.random() < 0.1 else None
            non_tax = round(rng.uniform(20, 200), 2) if rng.random() < 0.05 else None
            mgmt = u["mgmt"] if rent else 0
            expected = round(rent - mgmt - (repairs or 0) - (non_tax or 0), 2)
            transfer = expected
            if u["key"] == arrears_unit and k >= n_months - 4:
                transfer = round(expected * 0.5, 2)  # tenant paying half
            elif rng.random() < 0.08:
                transfer = round(expected - rng.choice([10, 50, 100]), 2)
            elec = round(rng.uniform(60, 220), 2) if u["kind"] == "mieszkanie" and rng.random() < 0.6 else None
            hoa = u["hoa"] if u["kind"] == "mieszkanie" and rng.random() < 0.7 else None
            fx_date = fx_rate = None
            tenant_payment = None
            if u["foreign"]:
                fx_date = dt.datetime(yy, mm, 15)
                fx_rate = round(rng.uniform(4.8, 5.3), 4)
                tenant_payment = round(rent * fx_rate, 2)
            elif rent and rng.random() < 0.4:
                tenant_payment = round(rent + (u["media"] or 0), 2)
            balance = round(transfer - expected, 2)
            taxable = round(transfer + (repairs or 0) + mgmt, 2)
            if u["key"] == units[3]["key"] and k == 2:
                taxable += 7  # a deliberate discrepancy for the data-quality check
            ws.append([None, u["key"], yy, u["manager"], u["city"], u["address"],
                       int(u["nr"]) if u["nr"].isdigit() else u["nr"], "Jan Kowalski", mm,
                       rent, u["media"] if rent else None, tenant_payment, mgmt, hoa, elec, repairs,
                       transfer, expected, balance, non_tax, taxable, fx_date, fx_rate])
            rows += 1
    ref = wb.create_sheet("Reference")
    ref.append(REFERENCE_HEADER)
    for u in units[:-1] if n_units > 3 else units:  # last unit only in the ledger
        ref.append([u["label"], u["key"], u["kind"], u["value_t0"], u["area"], u["floor"], None, dt.datetime.combine(u["t0"], dt.time())])
    ref.append(["Extra, not rented", unit_key("Krakow", "Nieznana 1", "9"), "mieszkanie", 400_000, 40.0, 2, None, dt.datetime(2020, 5, 5)])
    tr = wb.create_sheet("Trades")
    tr.append(TRADES_HEADER)
    day = dt.datetime(months[0][0], months[0][1], 3)
    tickers = ["ETF1.WA", "ACME", "BETA"]
    held = {t: 0.0 for t in tickers}
    for _ in range(min(40, n_months * 3)):
        t = rng.choice(tickers)
        side = "SELL" if held[t] > 0 and rng.random() < 0.35 else "BUY"
        qty = rng.randint(1, 20) if side == "BUY" else rng.randint(1, int(held[t]))
        held[t] += qty if side == "BUY" else -qty
        tr.append([day, t, side, qty, round(rng.uniform(20, 200), 2), round(rng.uniform(1, 9), 2), "PLN", "IKE"])
        day += dt.timedelta(days=rng.randint(3, 15))
    return wb, rows


def main(argv=None):
    ap = argparse.ArgumentParser(description="write a synthetic ledger workbook")
    ap.add_argument("out")
    ap.add_argument("--units", type=int, default=12)
    ap.add_argument("--months", type=int, default=18)
    ap.add_argument("--seed", type=int, default=7)
    a = ap.parse_args(argv)
    wb, rows = make_workbook(a.units, a.months, a.seed)
    wb.save(a.out)
    print(f"wrote {a.out}: {a.units} units, {rows} ledger rows")


if __name__ == "__main__":
    main()
