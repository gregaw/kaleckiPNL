"""xlsx bytes -> normalized polars frames. Pure: no I/O beyond the bytes handed in.

The ledger tab has one row per unit per month. Header row is found by content (it contains
'Lokal' and 'Rok'), not by position, and columns are matched by normalized prefix so wording
changes in the sheet do not break the import. Formula strings that Google exported without a
cached value (e.g. a custom NBP function) become null."""
from __future__ import annotations

import datetime as dt
import io
import re
import unicodedata
from dataclasses import dataclass, field

import openpyxl
import polars as pl

# Longest prefix first, so 'przelew spodziewany' wins over 'przelew'.
LEDGER_COLUMNS: list[tuple[str, str]] = [
    ("przelew spodziewany", "expected_transfer_sheet"),
    ("faktura za zarz", "mgmt_invoice"),
    ("faktura pge", "electricity"),
    ("oplata do wspolnoty", "hoa_fee"),
    ("naprawy", "repairs"),
    ("koszty poza podatkiem", "non_tax_costs"),
    ("zysk-do-podatku", "taxable_profit_sheet"),
    ("zysk do podatku", "taxable_profit_sheet"),
    ("zaliczka na media", "media_advance"),
    ("czynsz", "contract_rent"),
    ("wplata najemcy", "tenant_payment"),
    ("nadplata", "balance_sheet"),
    ("data kursu", "fx_date"),
    ("kurs nbp", "fx_rate"),
    ("kurs", "fx_rate"),
    ("przelew", "transfer"),
    ("lokal", "unit_key"),
    ("rok", "year"),
    ("zarzad", "manager"),
    ("miasto", "city"),
    ("adres", "address"),
    ("nr", "unit_no"),
    ("wl", "owner"),
    ("mc", "month"),
]
REFERENCE_COLUMNS: list[tuple[str, str]] = [
    ("mieszkanie", "unit_key"),
    ("typ", "kind"),
    ("wartosc", "value_t0"),
    ("powierzchnia", "area_m2"),
    ("pietro", "floor"),
    ("kw", "kw"),
    ("t0", "acquired_on"),
]
TRADE_COLUMNS: list[tuple[str, str]] = [
    ("date", "ts"), ("data", "ts"), ("ticker", "ticker"), ("symbol", "ticker"),
    ("side", "side"), ("qty", "qty"), ("quantity", "qty"), ("ilosc", "qty"),
    ("price", "price"), ("cena", "price"), ("fee", "fee"), ("prowizja", "fee"),
    ("currency", "currency"), ("waluta", "currency"), ("account", "account"), ("konto", "account"),
]
NUMERIC = ["contract_rent", "media_advance", "tenant_payment", "mgmt_invoice", "hoa_fee",
           "electricity", "repairs", "transfer", "expected_transfer_sheet", "balance_sheet",
           "non_tax_costs", "taxable_profit_sheet", "fx_rate"]
LEDGER_SCHEMA = {
    "unit_key": pl.Utf8, "year": pl.Int32, "month": pl.Int32, "manager": pl.Utf8,
    "city": pl.Utf8, "address": pl.Utf8, "unit_no": pl.Utf8, "owner": pl.Utf8,
    **{c: pl.Float64 for c in NUMERIC}, "fx_date": pl.Date, "row_no": pl.Int32,
}
UNITS_SCHEMA = {
    "unit_key": pl.Utf8, "label": pl.Utf8, "kind": pl.Utf8, "city": pl.Utf8, "address": pl.Utf8,
    "unit_no": pl.Utf8, "owner": pl.Utf8, "manager": pl.Utf8, "value_t0": pl.Float64,
    "area_m2": pl.Float64, "floor": pl.Float64, "kw": pl.Utf8, "acquired_on": pl.Date,
    "currency": pl.Utf8, "in_reference": pl.Boolean, "in_ledger": pl.Boolean,
}
TRADES_SCHEMA = {
    "trade_id": pl.Int32, "ts": pl.Datetime("us"), "ticker": pl.Utf8, "side": pl.Utf8,
    "qty": pl.Float64, "price": pl.Float64, "fee": pl.Float64, "currency": pl.Utf8,
    "account": pl.Utf8,
}


def norm(s) -> str:
    """lowercase, ascii-fold (ł -> l), collapse whitespace."""
    if s is None:
        return ""
    s = str(s).replace("ł", "l").replace("Ł", "L")
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).strip().lower()


def _map_header(row, rules) -> dict[int, str]:
    out: dict[int, str] = {}
    used: set[str] = set()
    for i, cell in enumerate(row):
        h = norm(cell)
        if not h:
            continue
        for prefix, name in rules:
            if h.startswith(prefix) and name not in used:
                out[i] = name
                used.add(name)
                break
    return out


def _find_header(rows, *must):
    for r, row in enumerate(rows):
        names = {norm(c) for c in row if c is not None}
        if all(any(n.startswith(m) for n in names) for m in must):
            return r
    return None


def _num(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        s = v.strip()
        if not s or s.startswith("=") or s in {"-", ".", "–"}:
            return None
        s = s.replace(" ", "").replace(" ", "").replace(",", ".")
        try:
            return float(s)
        except ValueError:
            return None
    return None


def _int(v):
    n = _num(v)
    return None if n is None else int(round(n))


def _text(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    s = str(v).strip()
    if s.startswith("="):
        return None
    return s or None


def _date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, str):
        for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y"):
            try:
                return dt.datetime.strptime(v.strip(), fmt).date()
            except ValueError:
                pass
    return None


def _datetime(v):
    if isinstance(v, dt.datetime):
        return v
    d = _date(v)
    return dt.datetime(d.year, d.month, d.day) if d else None


@dataclass
class ParseResult:
    ledger: pl.DataFrame
    units: pl.DataFrame
    trades: pl.DataFrame
    warnings: list[str] = field(default_factory=list)


def _sheet_rows(wb, name: str):
    if name not in wb.sheetnames:
        return None
    return [tuple(r) for r in wb[name].iter_rows(values_only=True)]


def parse_ledger(rows) -> tuple[pl.DataFrame, list[str]]:
    warnings: list[str] = []
    h = _find_header(rows, "lokal", "rok")
    if h is None:
        raise ValueError("ledger tab: no header row containing 'Lokal' and 'Rok'")
    cols = _map_header(rows[h], LEDGER_COLUMNS)
    missing = {"unit_key", "year", "month"} - set(cols.values())
    if missing:
        raise ValueError(f"ledger tab: missing columns {sorted(missing)}")
    recs = []
    for r in range(h + 1, len(rows)):
        row = rows[r]
        rec = {k: None for k in LEDGER_SCHEMA}
        for i, name in cols.items():
            v = row[i] if i < len(row) else None
            if name in NUMERIC:
                rec[name] = _num(v)
            elif name in ("year", "month"):
                rec[name] = _int(v)
            elif name == "fx_date":
                rec[name] = _date(v)
            else:
                rec[name] = _text(v)
        if not rec["unit_key"]:
            continue
        if rec["year"] is None or rec["month"] is None:
            warnings.append(f"ledger row {r + 1}: missing year/month, skipped")
            continue
        rec["row_no"] = r + 1
        recs.append(rec)
    df = pl.DataFrame(recs, schema=LEDGER_SCHEMA, orient="row") if recs else pl.DataFrame(schema=LEDGER_SCHEMA)
    return df.sort(["unit_key", "year", "month"]), warnings


def parse_reference(rows) -> pl.DataFrame:
    cols_out = ["unit_key", "label", "kind", "value_t0", "area_m2", "floor", "kw", "acquired_on"]
    schema = {k: UNITS_SCHEMA[k] for k in cols_out}
    if rows is None:
        return pl.DataFrame(schema=schema)
    h = _find_header(rows, "mieszkanie", "typ")
    if h is None:
        return pl.DataFrame(schema=schema)
    cols = _map_header(rows[h], REFERENCE_COLUMNS)
    key_idx = next(i for i, n in cols.items() if n == "unit_key")
    recs = []
    for row in rows[h + 1:]:
        rec = {k: None for k in cols_out}
        for i, name in cols.items():
            v = row[i] if i < len(row) else None
            if name in ("value_t0", "area_m2", "floor"):
                rec[name] = _num(v)
            elif name == "acquired_on":
                rec[name] = _date(v)
            else:
                rec[name] = _text(v)
        # the unlabeled first column holds a human label
        if key_idx > 0:
            rec["label"] = _text(row[0]) if len(row) else None
        if rec["unit_key"]:
            if rec["kind"]:
                rec["kind"] = norm(rec["kind"])
            recs.append(rec)
    return pl.DataFrame(recs, schema=schema, orient="row") if recs else pl.DataFrame(schema=schema)


def parse_trades(rows) -> pl.DataFrame:
    if rows is None:
        return pl.DataFrame(schema=TRADES_SCHEMA)
    h = _find_header(rows, "ticker", "side")
    if h is None:
        h = _find_header(rows, "symbol", "side")
    if h is None:
        return pl.DataFrame(schema=TRADES_SCHEMA)
    cols = _map_header(rows[h], TRADE_COLUMNS)
    recs = []
    for n, row in enumerate(rows[h + 1:], start=1):
        rec = {k: None for k in TRADES_SCHEMA}
        for i, name in cols.items():
            v = row[i] if i < len(row) else None
            if name in ("qty", "price", "fee"):
                rec[name] = _num(v)
            elif name == "ts":
                rec[name] = _datetime(v)
            else:
                rec[name] = _text(v)
        if not rec["ticker"] or rec["ts"] is None or rec["qty"] is None or rec["price"] is None:
            continue
        rec["trade_id"] = n
        rec["side"] = (rec["side"] or "BUY").strip().upper()
        rec["fee"] = rec["fee"] or 0.0
        rec["currency"] = (rec["currency"] or "PLN").upper()
        recs.append(rec)
    df = pl.DataFrame(recs, schema=TRADES_SCHEMA, orient="row") if recs else pl.DataFrame(schema=TRADES_SCHEMA)
    return df.sort(["ts", "trade_id"])


def build_units(ledger: pl.DataFrame, reference: pl.DataFrame, foreign_units: dict[str, str] | None = None) -> pl.DataFrame:
    """Union of the Reference tab and the units seen in the ledger. Descriptive fields come
    from the latest ledger row; currency is GBP where the sheet carries an NBP rate."""
    foreign_units = foreign_units or {}
    if ledger.height:
        latest = (ledger.sort(["year", "month"], descending=True)
                  .group_by("unit_key").agg([
                      pl.col("city").drop_nulls().first(),
                      pl.col("address").drop_nulls().first(),
                      pl.col("unit_no").drop_nulls().first(),
                      pl.col("owner").drop_nulls().first(),
                      pl.col("manager").drop_nulls().first(),
                      (pl.col("fx_rate").is_not_null().any() | pl.col("fx_date").is_not_null().any()).alias("has_fx"),
                  ]).with_columns(pl.lit(True).alias("in_ledger")))
    else:
        latest = pl.DataFrame(schema={"unit_key": pl.Utf8, "city": pl.Utf8, "address": pl.Utf8,
                                      "unit_no": pl.Utf8, "owner": pl.Utf8, "manager": pl.Utf8,
                                      "has_fx": pl.Boolean, "in_ledger": pl.Boolean})
    ref = reference.with_columns(pl.lit(True).alias("in_reference"))
    units = ref.join(latest, on="unit_key", how="full", coalesce=True)
    units = units.with_columns([
        pl.col("in_reference").fill_null(False),
        pl.col("in_ledger").fill_null(False),
        pl.col("has_fx").fill_null(False),
        pl.col("kind").fill_null("mieszkanie"),
    ])
    override = pl.DataFrame({"unit_key": list(foreign_units), "ccy_override": list(foreign_units.values())},
                            schema={"unit_key": pl.Utf8, "ccy_override": pl.Utf8})
    units = units.join(override, on="unit_key", how="left").with_columns(
        pl.coalesce([pl.col("ccy_override"), pl.when(pl.col("has_fx")).then(pl.lit("GBP")).otherwise(pl.lit("PLN"))]).alias("currency")
    ).drop(["ccy_override", "has_fx"])
    units = units.with_columns(pl.coalesce([pl.col("label"), pl.col("unit_key")]).alias("label"))
    return units.select(list(UNITS_SCHEMA)).sort("unit_key")


def parse_workbook(data: bytes, ledger_tab="All-Data", reference_tab="Reference",
                   trades_tab="Trades", foreign_units: dict[str, str] | None = None) -> ParseResult:
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        ledger_rows = _sheet_rows(wb, ledger_tab)
        if ledger_rows is None:
            raise ValueError(f"tab {ledger_tab!r} not found; tabs are {wb.sheetnames}")
        ledger, warnings = parse_ledger(ledger_rows)
        ref_rows = _sheet_rows(wb, reference_tab)
        if ref_rows is None:
            warnings.append(f"tab {reference_tab!r} not found; units come from the ledger only")
        reference = parse_reference(ref_rows)
        trades = parse_trades(_sheet_rows(wb, trades_tab))
    finally:
        wb.close()
    units = build_units(ledger, reference, foreign_units)
    return ParseResult(ledger=ledger, units=units, trades=trades, warnings=warnings)
